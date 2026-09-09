#!/usr/bin/env python3
"""
Build the fillable instrument forms for the Hospital Digitalisation Handbook.

Produces, in docs/forms/:
  - one .docx per instrument, all ten (fill in on a computer)
  - one .pdf per instrument (print and fill in by hand), converted with LibreOffice
  - D1_Data_Capture_KoBo.xlsx — an XLSForm to upload to KoBoToolbox (phone, offline, photos)

Run from the repository root:   python scripts/build_forms.py
Requires: python-docx, openpyxl, and LibreOffice (soffice) on PATH for the PDF step.

Keep this script and the handbook version in step: bump VERSION below when the
instruments change, and the forms carry it in their footer.
"""
import os, sys, subprocess, shutil
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

VERSION = "0.4"
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "forms")
os.makedirs(OUT, exist_ok=True)

NAVY = RGBColor(0x0B, 0x35, 0x83)
GREY = RGBColor(0x5A, 0x66, 0x72)

# ---------------------------------------------------------------- helpers ---
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, size=9.5, color=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(2)

def new_doc(code, title, stage, purpose, when):
    d = Document()
    for s in d.sections:
        s.page_height = Cm(29.7); s.page_width = Cm(21.0)
        s.left_margin = s.right_margin = Cm(1.8); s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.6)
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)

    # header block
    p = d.add_paragraph(); r = p.add_run(f"AUCOOP · Hospital Digitalisation Handbook · Instrument {code}")
    r.font.size = Pt(8.5); r.font.color.rgb = NAVY; r.bold = True
    p.paragraph_format.space_after = Pt(0)
    h = d.add_paragraph(); r = h.add_run(title); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = NAVY
    h.paragraph_format.space_after = Pt(2)
    m = d.add_paragraph(); r = m.add_run(f"{stage}   ·   {when}"); r.font.size = Pt(9); r.font.color.rgb = GREY
    m.paragraph_format.space_after = Pt(6)
    pp = d.add_paragraph(); r = pp.add_run(purpose); r.italic = True; r.font.size = Pt(9.5)
    pp.paragraph_format.space_after = Pt(8)

    # identity strip
    t = d.add_table(rows=1, cols=4); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, lab in enumerate(["Hospital", "Unit / area", "Completed by", "Date"]):
        set_cell_text(t.rows[0].cells[i], lab + "\n\n", bold=True, size=8.5, color=GREY)
        shade(t.rows[0].cells[i], "EEF2F7")
    d.add_paragraph().paragraph_format.space_after = Pt(4)

    # footer
    for s in d.sections:
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run(f"Handbook v{VERSION}  ·  {code}  ·  Record the handbook version in your closure report")
        fr.font.size = Pt(8); fr.font.color.rgb = GREY
    return d

def section(d, text):
    p = d.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = NAVY
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)

def note(d, text):
    p = d.add_paragraph(); r = p.add_run(text); r.font.size = Pt(8.5); r.font.color.rgb = GREY; r.italic = True
    p.paragraph_format.space_after = Pt(4)

def table(d, header, rows, widths=None, blank_rows=0, blank_cols=None):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        set_cell_text(t.rows[0].cells[i], h, bold=True, size=8.5, color=GREY); shade(t.rows[0].cells[i], "EEF2F7")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell_text(cells[i], v, size=9)
    for _ in range(blank_rows):
        cells = t.add_row().cells
        for i in range(len(header)):
            set_cell_text(cells[i], "\n", size=9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def checklist(d, items):
    for it in items:
        p = d.add_paragraph(); r = p.add_run("☐  "); r.font.size = Pt(11)
        if isinstance(it, tuple):
            r2 = p.add_run(it[0]); r2.font.size = Pt(9.5)
            r3 = p.add_run("  — " + it[1]); r3.font.size = Pt(8.5); r3.font.color.rgb = GREY
        else:
            r2 = p.add_run(it); r2.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(2)

def lines(d, label, n=3):
    p = d.add_paragraph(); r = p.add_run(label); r.bold = True; r.font.size = Pt(9.5)
    p.paragraph_format.space_after = Pt(1)
    for _ in range(n):
        q = d.add_paragraph("_" * 92); q.paragraph_format.space_after = Pt(1)
        for rr in q.runs: rr.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB); rr.font.size = Pt(9)
    d.add_paragraph().paragraph_format.space_after = Pt(1)

# ---------------------------------------------------------------- forms ---
def build_B1():
    d = new_doc("B-1", "Readiness assessment", "Before you go",
        "Score the gap between what the system will assume and what the hospital actually is. Higher total = higher risk. Feeds B-2.",
        "Complete remotely before travel; revise on site in week one")
    section(d, "1. Design–reality gap — score each dimension 0 (no gap) to 10 (total gap)")
    note(d, "After Heeks' design–reality gap model. Rate the distance between what the proposed system assumes and what you find.")
    table(d, ["Dimension", "What you are rating", "Score 0–10", "Evidence / note"], [
        ["Information", "What the system expects recorded vs. what is written down today", "", ""],
        ["Technology", "Power, network, hardware assumed vs. installed and working", "", ""],
        ["Processes", "The workflow the software implies vs. the real patient journey (D-3)", "", ""],
        ["Objectives & values", "What the deploying team wants vs. what the hospital wants", "", ""],
        ["Staffing & skills", "Digital literacy assumed vs. present on the ward", "", ""],
        ["Management systems", "Governance and accountability assumed vs. actual", "", ""],
        ["Other resources", "Money and time assumed vs. committed", "", ""],
        ["TOTAL", "", "", ""],
    ], widths=[3.4, 7.2, 2.0, 4.8])
    section(d, "2. The four questions that decide survival")
    table(d, ["Code", "Question", "Answer", "Detail (name, figure, source)"], [
        ["PAID-IT", "Is there a paid IT person at the hospital?", "Yes / No", ""],
        ["OWNER-5Y", "Is there a funded owner for the system for the next five years?", "Yes / No", ""],
        ["PRESENCE", "How many years is the organisation committed to this site?", "", ""],
        ["HANDOVER", "Will a handover pack (L-1) exist on departure?", "Yes / No", ""],
    ], widths=[2.4, 7.6, 2.2, 5.2])
    section(d, "3. Result")
    note(d, "Thresholds and scale anchors are not yet calibrated (v0.4). Use the score to structure the conversation, then complete B-2.")
    lines(d, "Summary judgement and who you discussed it with:", 4)
    return d

def build_B2():
    d = new_doc("B-2", "Level decision", "Before you go",
        "Decide how far up the tool spectrum this deployment goes — driven by what the hospital can sustain after you leave, not by need.",
        "Complete after B-1, before any equipment is bought")
    section(d, "1. The stop rule — tick every statement that is TRUE")
    note(d, "If ANY box is ticked, drop a level or do not deploy. Dropping a level is not a defeat.")
    checklist(d, [
        ("No named person is paid to keep the system running after we leave", "not 'someone will look after it' — a name, a role, a salary line"),
        ("The medical director has not personally endorsed the project", "an administrator's enthusiasm is not the same authority"),
        ("There is no route to fund maintenance for five years", "hardware fails, software needs patching, someone must be paid"),
        ("The unit we intend to digitise has no verified working power AND network", "verified by us, in the room, with a device"),
        ("We cannot retire the paper process for what we are deploying", "if paper continues alongside, nothing has been digitised"),
    ])
    section(d, "2. The tool spectrum")
    table(d, ["Level", "Scope", "Tools", "Minimum to keep it alive"], [
        ["1", "Data collection & surveys", "KoBoToolbox, ODK", "A trained clerk; no server on a hosted instance"],
        ["2", "Aggregate reporting & case tracking", "DHIS2, CommCare", "A reporting officer; alignment with the national system"],
        ["3", "Clinical EMR", "OpenMRS", "A part-time IT person; reliable power where clinicians work"],
        ["4", "Full HIS — clinical, lab, pharmacy, billing", "Bahmni, GNU Health", "A paid full-time IT person, a funded five-year owner, a support arrangement"],
    ], widths=[1.4, 5.0, 3.6, 7.4])
    section(d, "3. Decision")
    table(d, ["Level chosen (1–4, or 'do not deploy')", "B-1 total score", "Decided by (names)", "Date"], [["", "", "", ""]], widths=[6.0, 3.0, 5.4, 3.0])
    lines(d, "Justification — why this level and not the one above:", 5)
    lines(d, "If we deploy below what was hoped: what has to become true for next year's answer to be higher?", 4)
    note(d, "That last line is a prediction with a date on it. Next year's team is asked whether it came true — so write something that can be checked.")

    section(d, "4. Stop-rule override — complete this ONLY if you ticked a box in section 1 and are deploying anyway")
    note(d, "This is the most important section on the form. In 2026 at Lunsar every stop-rule condition was true and a level-4 system was deployed by careful engineers with reasons that were never written down. An override is sometimes right. An unrecorded override is how a decision becomes invisible to everyone who comes after.")
    table(d, ["Which conditions were true", "Deploying anyway?", "Decided by (names, roles)", "Date"], [
        ["", "Yes / No", "", ""],
    ], widths=[6.0, 2.6, 5.4, 3.4])
    lines(d, "On what grounds — the actual reasoning, not the conclusion:", 4)
    lines(d, "What is being put in place to compensate, and who is accountable for it:", 4)
    lines(d, "What would have to be true for us to reverse this decision, and who reviews it:", 3)
    return d

def build_B3():
    d = new_doc("B-3", "Mandate and ownership", "Before you go",
        "Name the six people whose support decides whether this works, speak to each separately, and put the agreement in writing.",
        "Before configuration starts")
    section(d, "1. The six people — spoken to separately, by name")
    table(d, ["Role", "Name", "Met on (date)", "Q1: What takes too long in your unit?", "Q2: What would make this worth your staff's time?"], [
        ["Medical Director", "", "", "", ""],
        ["Head of Nursing", "", "", "", ""],
        ["Head of Pharmacy", "", "", "", ""],
        ["Laboratory Director", "", "", "", ""],
        ["Administrator / Finance", "", "", "", ""],
        ["IT Officer / technician", "", "", "", ""],
    ], widths=[3.2, 3.0, 2.2, 4.5, 4.5])
    note(d, "If the IT Officer post does not exist, go back to B-2 — that is a stop-rule condition.")
    section(d, "1b. Permission or instruction?")
    note(d, "A director will say yes to visiting engineers. That yes changes nothing on the wards. The mandate is real only when the director tells staff, personally, that this is how the hospital will now work. This is the field that separates the sites that adopted from the sites that did not.")
    table(d, ["Did the director announce it to staff?", "To whom", "When", "In what form (meeting / memo / word to each head)"], [
        ["Yes / No / Promised, not yet done", "", "", ""],
    ], widths=[4.4, 3.6, 2.6, 6.8])
    note(d, "If the answer is No, or if the director expects YOU to do the telling, you have permission and not a mandate. Say so here rather than proceeding as though you had one.")
    section(d, "2. Super-user candidates (permanent staff, respected, curious — not the most senior)")
    table(d, ["Name", "Unit", "Why them", "Time allowance agreed?"], [["", "", "", ""], ["", "", "", ""]], widths=[4.0, 3.0, 7.0, 3.4])
    section(d, "3. The written agreement (MoU) — tick when the clause exists in writing")
    checklist(d, [
        "Who owns the system — the hospital, stated explicitly",
        "Who owns the data, and how patient consent is handled",
        "Training commitments: how many staff, released for how long, when",
        "Remote support: what the organisation will and will not do after leaving, and for how long",
        "Who pays for what after we leave: hosting, hardware replacement, the IT post",
        "Signed by the medical director",
    ])
    lines(d, "Existing systems discovered (finance package, any previous digital attempt, who ran it):", 4)
    return d

def build_D1():
    d = new_doc("D-1", "Field data capture — paper sheet", "While you are there",
        "Paper fallback for the KoBoToolbox form. One sheet per unit. Photograph every form in use — one blank, one completed — before you leave the building.",
        "Every unit, before leaving each building")
    note(d, "Prefer the KoBoToolbox form (D1_Data_Capture_KoBo.xlsx) on a phone: it attaches photos directly and works offline. Use this sheet if no phone is available.")
    section(d, "1. Paper forms in use in this unit")
    table(d, ["Form name (as printed)", "Blank photo taken?", "Completed photo taken?", "Who fills it", "Who receives it, by when", "Per day / month"], [], widths=[4.2, 1.8, 1.9, 2.8, 3.6, 2.2], blank_rows=8)
    section(d, "2. Volume — counted, not estimated")
    table(d, ["Measure", "Count", "How counted (register / observed / asked)", "Date"], [
        ["Patients seen per day", "", "", ""], ["Inpatients today", "", "", ""],
        ["Lab tests per day", "", "", ""], ["Deliveries per month", "", "", ""], ["Other:", "", "", ""],
    ], widths=[4.0, 2.4, 7.2, 3.0])
    section(d, "3. How is the patient identified at this point?")
    lines(d, "Identifier used (folder ID? name? other?), where it is written, whether this unit issues its own numbers:", 3)
    section(d, "4. Power and network in this room")
    table(d, ["Item", "Reading / answer"], [
        ["Working socket in the room? How many? Which circuit?", ""],
        ["Grid voltage measured (V) — and time of reading", ""],
        ["Grid hours per day (observed)", ""],
        ["Working network link IN THIS ROOM? (yes / no / intermittent)", ""],
        ["Latency measured from here with the real application (ms)", ""],
    ], widths=[8.0, 8.6])
    section(d, "5. Room photograph taken? (outlets, cable routes, where a workstation could go)")
    checklist(d, ["Yes — file name(s):"])
    lines(d, "Notes — anything the next team could not learn from Barcelona:", 5)
    return d

def build_D2():
    d = new_doc("D-2", "Infrastructure audit", "While you are there",
        "The inputs for power, network and server dimensioning. Completed during the assessment visit.",
        "Assessment visit, first week")
    for cat, items in [
        ("Electricity", ["Grid availability (hours/day, observed)", "Generator: fuel type / capacity / runtime / fuel reliability", "Solar installation: panels, rating, orientation, battery capacity & age", "Voltage stability — 24 h multimeter readings (min / max / typical)", "Outlet locations per building"]),
        ("Internet", ["ISP name and technology (fibre / 4G / VSAT / Starlink)", "Contracted bandwidth", "Measured bandwidth (down / up, time of day)", "Carrier-grade NAT check: router WAN address vs public address — same or different?", "Monthly cost"]),
        ("Physical space", ["Candidate server room: ventilation / security / access control / who has a key", "Cable routes between buildings; anything crossing open ground or a road", "Distance between the farthest buildings", "Wall materials"]),
        ("Existing IT", ["Functional computers and tablets (count, OS, age)", "Printers", "Existing network equipment (models, where, working?)", "Existing cabling", "Existing software systems (finance, forms, anything from a previous project)"]),
        ("Environment", ["Temperature range", "Humidity", "Dust", "Insect / rodent risk", "Flooding risk", "Lightning — frequency, existing protection"]),
    ]:
        section(d, cat)
        table(d, ["Item", "Finding", "Photo / file"], [[i, "", ""] for i in items], widths=[7.4, 6.6, 2.6])
    return d

def build_D3():
    d = new_doc("D-3", "Workflow mapping", "While you are there",
        "Document the patient journey as it actually is, before proposing to change it. Confirm or correct each step on site.",
        "First week, before configuring anything")
    section(d, "1. The patient journey — confirm, correct, or strike out each step")
    table(d, ["#", "Step", "Current method (register / logbook / verbal)", "Daily volume", "Staff involved", "Pain point THEY name"], [
        ["1", "Registration", "", "", "", ""], ["2", "Triage", "", "", "", ""], ["3", "Consultation", "", "", "", ""],
        ["4", "Laboratory", "", "", "", ""], ["5", "Prescription", "", "", "", ""], ["6", "Pharmacy", "", "", "", ""],
        ["7", "Billing", "", "", "", ""], ["8", "Discharge / admission", "", "", "", ""], ["", "Other:", "", "", "", ""],
    ], widths=[0.8, 2.8, 4.0, 2.0, 3.0, 4.0])
    section(d, "2. The copy chain")
    note(d, "Every place a figure is copied from one book into another is where errors enter and time is lost — and where a digital system earns its keep.")
    lines(d, "Draw or describe the chain from 'a patient walks in' to 'a number appears in a monthly return':", 8)
    section(d, "3. The one process that could become digital-only")
    lines(d, "Which single process could retire its paper register within this deployment, and what has to be true first?", 4)
    return d

def build_D4():
    d = new_doc("D-4", "Phase exit gate", "While you are there",
        "A phase is not finished because the feature works. Complete one sheet per phase you claim to have closed.",
        "At the end of each phase")
    section(d, "Phase being closed")
    table(d, ["Phase (1–4)", "Name", "Claimed closed on", "By"], [["", "", "", ""]], widths=[2.4, 6.0, 4.0, 4.2])
    section(d, "1. Engineering conditions — evidence for each")
    table(d, ["Condition", "Met? (Y/N)", "Evidence (what was shown, to whom)"], [
        ["Phase-specific gate (see handbook D-4 table)", "", ""],
        ["Works offline-tolerantly — degrades sensibly when power or network fails", "", ""],
        ["A hospital-usable runbook exists, written for zero prior IT experience", "", ""],
        ["The runbook was tested on someone who was NOT in the training", "", ""],
    ], widths=[8.0, 2.2, 6.4])
    section(d, "2. The four measured outcomes")
    note(d, "Use these definitions exactly \u2014 S-1 re-measures the same four at six and twelve months, and a measure whose definition changes between cohorts cannot be compared across them. Write the denominator, not just the ratio.")
    table(d, ["Outcome", "Definition \u2014 use exactly this", "Numerator", "Denominator", "Value"], [
        ["Adoption", "Eligible staff who used it in a normal week \u00f7 eligible staff", "", "", ""],
        ["Fidelity", "Records entered at the point of care \u00f7 records entered (sample one real day)", "", "", ""],
        ["Penetration", "Units live \u00f7 units in scope at handover (name them)", "", "", ""],
        ["Sustainability", "Running, with a named owner in post, on this date", "", "", ""],
    ], widths=[2.8, 6.6, 2.4, 2.4, 2.2])

    section(d, "3. The other four \u2014 one sentence each, in prose")
    note(d, "These will not be measured honestly in the last week of a three-week deployment. A form that returns three real numbers and five guesses produces something that looks like data. Write a sentence.")
    table(d, ["Outcome", "The question", "What you observed"], [
        ["Acceptability", "Do staff find it agreeable?", ""],
        ["Appropriateness", "Does it fit this setting?", ""],
        ["Feasibility", "Can it be done here, with these people?", ""],
        ["Implementation cost", "Hardware, travel, staff time, recurrent", ""],
    ], widths=[3.0, 5.4, 9.0])

    section(d, "4. Six- and twelve-month check")
    note(d, "Sustainability is not measurable at departure \u2014 it is measurable in a year. Name the person and the dates here, and again in L-1 section 4. They will be filling in S-1.")
    table(d, ["Who checks", "6-month date", "12-month date", "Contact route that will still work then"], [["", "", "", ""]], widths=[4.0, 3.0, 3.0, 7.4])
    section(d, "5. What was not done")
    lines(d, "What was NOT done in this phase, and why (the next team's scope starts here):", 5)
    return d

def build_L1():
    d = new_doc("L-1", "Handover pack — cover sheet", "Before you leave",
        "What physically stays at the hospital. Started mid-second-week, walked through with the named owner in person before departure.",
        "Started on Wednesday of week two; complete before the last day")
    section(d, "1. Named owner")
    table(d, ["Name", "Role", "Contact", "Walked through the pack on (date)"], [["", "", "", ""]], widths=[4.4, 4.0, 4.2, 4.0])
    section(d, "2. Contents — tick when present and handed over")
    checklist(d, [
        ("Runbook per deployed unit", "daily task, screenshots, written for zero prior IT; tested on someone not in the training"),
        ("'What to do when it breaks' sheet", "the five most likely failures and the first thing to try; includes 'the system takes 3–12 minutes to start'"),
        ("Credentials handed over securely", "in person, separately, recorded for rotation — NEVER written in this pack"),
        ("Support escalation ladder", "tier 0 in-product help · tier 1 on-site super-user · tier 2 external — each with a name and route"),
        ("Backup demonstrated by the owner while we watched", "backup AND restore"),
        ("Cutover plan, if paper is retired after departure", "date, who announces it, what happens on the day"),
        ("Who checks it at 6 and 12 months", "name and date"),
        ("What we deliberately did NOT do, and why", "the most useful page in the pack"),
        ("The handbook version we followed", f"v{VERSION}"),
    ])
    section(d, "3. Support ladder")
    table(d, ["Tier", "What", "Name", "Route (phone / email / where)"], [
        ["0", "Self-help in the product", "—", ""], ["1", "On-site super-user", "", ""], ["2", "External / commercial support", "", ""],
    ], widths=[1.4, 5.0, 4.0, 6.2])
    section(d, "4. Six- and twelve-month check")
    table(d, ["Check", "Who", "Date planned", "Done? Result"], [["6 months", "", "", ""], ["12 months", "", "", ""]], widths=[3.0, 4.6, 3.6, 5.4])
    lines(d, "What we did not do, and why:", 8)
    return d

def build_A1():
    d = new_doc("A-1", "Debrief and feedback return", "Once you are back",
        "What goes back into the handbook. Started in the field, closed within two weeks. Under two hours. Attach to the CCD closure report.",
        "Started on site; closed within two weeks of return")
    section(d, "0. What we inherited")
    note(d, "The handbook exists because each cohort has started from zero. This section is how that stops being a claim about the past and becomes something measured every year. Answer it even if — especially if — the answer is 'nothing'.")
    table(d, ["Question", "Answer", "Detail"], [
        ["Was there a handover pack (L-1) from a previous team?", "Yes / No / Partial", ""],
        ["Did you find completed instruments from a previous deployment?", "Yes / No — which?", ""],
        ["Were they usable — could you act on them without re-doing the work?", "Yes / No / Partly", ""],
        ["Was there a named person at the hospital who could explain what existed?", "Yes / No", ""],
        ["Did a six- or twelve-month check (S-1) exist for the previous deployment?", "Yes / No", ""],
    ], widths=[7.4, 3.6, 6.4])
    lines(d, "What was missing that cost you time — be specific, and estimate the days it cost:", 4)
    lines(d, "What you had to rediscover that someone already knew:", 3)

    section(d, "1. Version and completion")
    table(d, ["Handbook version followed", "Trip dates", "Site", "Team"], [["", "", "", ""]], widths=[4.0, 4.0, 4.6, 4.0])
    section(d, "2. Instruments — completed, skipped, and why")
    note(d, "A skipped instrument is a finding about the handbook, not a failure of the team. Say so plainly.")
    table(d, ["Instrument", "Completed?", "If skipped or partial — why?"], [[c, "", ""] for c in ["B-1", "B-2", "B-3", "D-1", "D-2", "D-3", "D-4", "L-1"]] + [["S-1", "Scheduled? Y / N", "Named owner and dates, from L-1 section 4 — this is the only outcome measurement the deployment will ever produce"]], widths=[2.4, 2.6, 11.6])
    section(d, "2b. Where you departed from the handbook, and why")
    note(d, "Different from section 3. A departure can be entirely correct and still needs recording — the next team's first question about any configuration is 'why is it set up this way?', and an unanswered why gets redesigned from scratch.")
    table(d, ["What the handbook says", "What we did instead", "Why", "Would you do it again?"], [], widths=[4.4, 4.4, 5.2, 3.4], blank_rows=5)

    section(d, "3. Where the handbook was wrong on site")
    lines(d, "Page or instrument, what it said, what was actually true:", 6)
    section(d, "4. What you had to invent because it was not covered")
    lines(d, "", 5)
    section(d, "5. Outcome scores at departure (from D-4)")
    table(d, ["Adoption", "Fidelity", "Penetration", "Sustainability (fill at 6 / 12 months)"], [["", "", "", ""]], widths=[4.0, 4.0, 4.0, 4.6])
    section(d, "6. The three questions you wish you had asked while still there")
    lines(d, "", 4)
    section(d, "7. Who can answer questions about this deployment in a year")
    note(d, "A student email address expires. Give a route that will still work when the next team reads this.")
    table(d, ["Name", "Role then", "Contact route with a real lifespan", "Until when"], [["", "", "", ""]], widths=[4.0, 3.6, 6.4, 3.4])
    lines(d, "The one thing the next team should do first:", 2)
    return d

def build_S1():
    d = new_doc("S-1", "Six- and twelve-month check", "Six and twelve months later",
        "Did it survive? The only instrument completed after the team has gone home — and the only one that measures an outcome rather than an intervention. Ten minutes, remotely, by the person named in L-1 section 4.",
        "At six months and again at twelve months, on the dates written in L-1")
    note(d, "Designed to be answerable from one conversation with the super-user, by phone or messaging. You are not auditing anyone. A blank field is a finding; an invented one is worse than nothing.")

    section(d, "1. Which check is this")
    table(d, ["Check", "Deployment it follows (site, dates)", "Handbook version that deployment followed", "Answered by, and how reached"], [
        ["\u2610 6 months   \u2610 12 months", "", "", ""],
    ], widths=[3.4, 5.0, 4.0, 5.0])

    section(d, "2. Is it running")
    table(d, ["Question", "Answer", "How you know"], [
        ["Is the system reachable / running at all?", "Yes / No / Partly", ""],
        ["Units live now, out of the units handed over", "     /     ", ""],
        ["Who is using it — roles and roughly how many people", "", ""],
        ["Is it used at the point of care, or entered later from paper?", "Point of care / Batched / Both", ""],
    ], widths=[7.0, 4.4, 6.0])

    section(d, "3. Paper")
    note(d, "The handbook's central design rule is that a process moves only when the digital path is the ONLY path. This is where you find out whether it stayed that way.")
    table(d, ["Process retired at handover (name it)", "Still retired?", "If it came back — when, and what triggered it"], [], widths=[5.6, 2.8, 9.0], blank_rows=3)

    section(d, "4. The owner")
    table(d, ["Question", "Answer"], [
        ["Is the named owner from L-1 still in post?", "Yes / No — if no, who now?"],
        ["Is the post funded, and by whom?", ""],
        ["Are the super-users still in place, and do they still have the time allowance?", ""],
        ["Has anyone at the hospital been paid, in the period, for work on this system?", ""],
    ], widths=[9.4, 8.0])

    section(d, "5. What broke since we left")
    note(d, "Including anything that broke and was NOT fixed. Silent downtime is the failure mode this project has repeatedly missed \u2014 a remote check in February 2026 found a server inactive and there is no record of what happened next.")
    table(d, ["What happened", "Roughly when", "Who dealt with it", "How long was it down?", "Did they need outside help?"], [], widths=[5.0, 2.6, 3.4, 3.0, 3.4], blank_rows=4)

    section(d, "6. B-1 re-scored \u2014 the same seven dimensions, today")
    note(d, "Score exactly as in B-1, 0 (no gap) to 10 (total gap), against the same anchors. The point is the movement, not the number: a gap that has not closed is the finding.")
    table(d, ["Dimension", "Score at B-1 (before travel)", "Score at departure (D-4)", "Score today", "What moved, or why nothing did"], [
        ["Information", "", "", "", ""], ["Technology", "", "", "", ""], ["Processes", "", "", "", ""],
        ["Objectives & values", "", "", "", ""], ["Staffing & skills", "", "", "", ""],
        ["Management systems", "", "", "", ""], ["Other resources", "", "", "", ""],
        ["TOTAL", "", "", "", ""],
    ], widths=[3.2, 2.6, 2.6, 2.2, 6.8])

    section(d, "7. The four outcomes, now")
    note(d, "The same four as D-4, measured the same way. If you cannot measure one from here, write 'not measurable remotely' \u2014 that is itself worth knowing about the measure.")
    table(d, ["Outcome", "Definition used", "Value at departure", "Value now"], [
        ["Adoption", "Eligible staff using it in a normal week \u00f7 eligible staff", "", ""],
        ["Fidelity", "Records entered at the point of care \u00f7 records entered", "", ""],
        ["Penetration", "Units live \u00f7 units in scope at handover", "", ""],
        ["Sustainability", "Running, with a named owner in post, on this date", "", ""],
    ], widths=[3.0, 7.4, 3.4, 3.6])

    section(d, "8. What the hospital wants that it is not getting")
    lines(d, "In their words. This is the next deployment's requirements list, gathered a year early and for free:", 4)

    section(d, "9. Where this goes")
    checklist(d, [
        ("Filed with the deployment's other completed instruments", "same folder, same naming convention"),
        ("Sent to whoever merges the A-1 returns", "if that role is still unnamed, send it to the handbook repository as an issue"),
        ("The next check is scheduled", "12-month check booked now, with a name and a date \u2014 or, at 12 months, closed"),
    ])
    table(d, ["Next check due", "Who will do it", "Contact route that will still work then"], [["", "", ""]], widths=[4.0, 5.4, 8.0])
    return d

# ---------------------------------------------------------------- XLSForm (KoBoToolbox) ---
def build_kobo_xlsform():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active; ws.title = "survey"
    ws.append(["type", "name", "label", "hint", "required", "appearance", "constraint", "relevant"])
    rows = [
        ["start", "start", "", "", "", "", "", ""],
        ["end", "end", "", "", "", "", "", ""],
        ["today", "today", "", "", "", "", "", ""],
        ["text", "hospital", "Hospital", "", "yes", "", "", ""],
        ["text", "unit", "Unit / area (e.g. OPD, Pharmacy, Emergency)", "", "yes", "", "", ""],
        ["text", "recorder", "Your name", "", "yes", "", "", ""],
        ["begin_group", "forms", "1. Paper forms in use in this unit", "One entry per form. Photograph a BLANK copy and a COMPLETED copy.", "", "field-list", "", ""],
        ["begin_repeat", "form", "Paper form", "", "", "", "", ""],
        ["text", "form_name", "Form name, as printed on it", "", "yes", "", "", ""],
        ["select_one form_type", "form_type", "What kind of form is it?", "", "yes", "", "", ""],
        ["image", "photo_blank", "Photo — BLANK form", "Whole sheet in frame, legible", "yes", "", "", ""],
        ["image", "photo_completed", "Photo — COMPLETED example", "Cover any patient names before photographing", "", "", "", ""],
        ["text", "filled_by", "Who fills it in (role)", "", "", "", "", ""],
        ["text", "goes_to", "Who receives it, and by when", "", "", "", "", ""],
        ["integer", "per_period", "How many per day (or per month for returns)", "", "", "", "", ""],
        ["select_one period", "period", "Per…", "", "", "", "", ""],
        ["end_repeat", "", "", "", "", "", "", ""],
        ["end_group", "", "", "", "", "", "", ""],
        ["begin_group", "volume", "2. Volume — counted, not estimated", "", "", "field-list", "", ""],
        ["integer", "opd_per_day", "Patients seen today (count)", "", "", "", "", ""],
        ["integer", "inpatients", "Inpatients today (count)", "", "", "", "", ""],
        ["integer", "lab_per_day", "Lab tests today (count)", "", "", "", "", ""],
        ["integer", "deliveries_month", "Deliveries last month (count)", "", "", "", "", ""],
        ["select_one counted_how", "counted_how", "How were these counted?", "", "", "", "", ""],
        ["end_group", "", "", "", "", "", "", ""],
        ["begin_group", "identity", "3. Patient identity at this point", "", "", "field-list", "", ""],
        ["select_multiple identifier", "identifier", "How is the patient identified here?", "", "yes", "", "", ""],
        ["select_one yesno", "own_numbers", "Does this unit issue its OWN patient numbers?", "If yes, that is an important finding — note it", "yes", "", "", ""],
        ["text", "identity_notes", "Notes on identification", "", "", "", "", ""],
        ["end_group", "", "", "", "", "", "", ""],
        ["begin_group", "power_net", "4. Power and network in this room", "", "", "field-list", "", ""],
        ["integer", "sockets", "Working sockets in the room (count)", "", "", "", "", ""],
        ["decimal", "voltage", "Grid voltage measured now (V)", "Multimeter reading", "", "", ". >= 100 and . <= 300", ""],
        ["integer", "grid_hours", "Grid hours per day, observed", "", "", "", ". >= 0 and . <= 24", ""],
        ["select_one link", "link", "Working network link IN THIS ROOM?", "Test with a device, here", "yes", "", "", ""],
        ["integer", "latency_ms", "Latency from here with the real application (ms)", "", "", "", "", "${link} != 'no'"],
        ["image", "photo_room", "Photo of the room — outlets, cable routes, where a workstation could go", "", "yes", "", "", ""],
        ["end_group", "", "", "", "", "", "", ""],
        ["text", "notes", "5. Anything the next team could not learn from Barcelona", "", "", "multiline", "", ""],
        ["geopoint", "location", "Location (optional)", "", "", "", "", ""],
    ]
    for r in rows: ws.append(r)

    ch = wb.create_sheet("choices")
    ch.append(["list_name", "name", "label"])
    for r in [
        ["form_type", "statutory", "Statutory return (to Ministry / association)"],
        ["form_type", "working", "Working record (register, logbook, day-book)"],
        ["form_type", "patient", "Patient-held (booklet, card)"],
        ["form_type", "other", "Other"],
        ["period", "day", "day"], ["period", "week", "week"], ["period", "month", "month"],
        ["counted_how", "register", "Counted from the register"], ["counted_how", "observed", "Observed / counted in person"], ["counted_how", "asked", "Asked a staff member (estimate)"],
        ["identifier", "folder_id", "Hospital folder / file number"], ["identifier", "name", "Name only"], ["identifier", "national_id", "National ID"], ["identifier", "booklet", "Patient-held booklet / card"], ["identifier", "other", "Other"],
        ["yesno", "yes", "Yes"], ["yesno", "no", "No"],
        ["link", "yes", "Yes — works"], ["link", "intermittent", "Intermittent"], ["link", "no", "No link"],
    ]: ch.append(r)

    st = wb.create_sheet("settings")
    st.append(["form_title", "form_id", "version", "default_language"])
    st.append(["D-1 Field data capture — Hospital Digitalisation Handbook", "hdh_d1_data_capture", VERSION, "English (en)"])
    for w in (ws, ch, st):
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = 34
    path = os.path.join(OUT, "D1_Data_Capture_KoBo.xlsx"); wb.save(path); return path

# ---------------------------------------------------------------- main ---
FORMS = {
    "B1_Readiness_Assessment": build_B1,
    "B2_Level_Decision": build_B2,
    "B3_Mandate_and_Ownership": build_B3,
    "D1_Data_Capture_Sheet": build_D1,
    "D2_Infrastructure_Audit": build_D2,
    "D3_Workflow_Mapping": build_D3,
    "D4_Phase_Exit_Gate": build_D4,
    "L1_Handover_Pack": build_L1,
    "A1_Debrief_and_Feedback": build_A1,
    "S1_Six_and_Twelve_Month_Check": build_S1,
}

def main():
    made = []
    for name, fn in FORMS.items():
        p = os.path.join(OUT, name + ".docx"); fn().save(p); made.append(p)
    made.append(build_kobo_xlsform())
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        docx = [p for p in made if p.endswith(".docx")]
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", OUT] + docx,
                       check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    else:
        print("LibreOffice not found — .docx written, PDFs skipped", file=sys.stderr)
    for f in sorted(os.listdir(OUT)):
        print(f"  {f:40s} {os.path.getsize(os.path.join(OUT, f))//1024:5d} KB")

if __name__ == "__main__":
    main()
